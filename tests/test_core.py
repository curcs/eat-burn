import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import profile_calc
import report
from llm import to_draft
from nutrition import Draft, Item, Nutrition, Per100
from parser import parse_edit_lines, parse_quick
from storage import Storage


# --- профиль ---

def test_target_for_owner():
    assert round(profile_calc.bmr(60, 165, 30)) == 1320
    assert profile_calc.target(60, 165, 30) == 1350


# --- разбор простых записей ---

@pytest.mark.parametrize("text,kind,kcal,desc", [
    ("450 рацион обед", "quick", 450, "рацион обед"),
    ("+300 ужин", "quick", 300, "ужин"),
    ("1 100 рацион", "quick", 1100, "рацион"),
    ("-200 бег", "workout", 200, "бег"),
    ("- 350 ккал силовая", "workout", 350, "силовая"),
    ("520", "quick", 520, "еда"),
])
def test_quick(text, kind, kcal, desc):
    q = parse_quick(text)
    assert (q.kind, q.kcal, q.desc) == (kind, kcal, desc)


@pytest.mark.parametrize("text", ["100г творога", "150 г курицы", "овсянка 60г", "3 шт печенья", "1 л кефира"])
def test_not_quick(text):
    assert parse_quick(text) is None


@pytest.mark.parametrize("text,kcal,desc", [("2 яйца", 2, "яйца"), ("4 л-карнитин сироп", 4, "л-карнитин сироп")])
def test_ambiguous_small_number(text, kcal, desc):
    q = parse_quick(text)
    assert (q.kind, q.kcal, q.desc) == ("ambiguous", kcal, desc)


def test_edit_lines():
    assert parse_edit_lines("банан 150, -мёд\nсыр 30г") == [
        ("банан", 150, False), ("мёд", None, True), ("сыр", 30, False)]


# --- поиск по базам ---

@pytest.fixture
def st(tmp_path):
    return Storage(str(tmp_path / "t.db"))


@pytest.fixture
def usda(tmp_path):
    p = tmp_path / "usda.db"
    db = sqlite3.connect(p)
    db.executescript("""
        CREATE TABLE foods (id INTEGER PRIMARY KEY, description TEXT, kcal REAL, protein REAL, fat REAL, carbs REAL);
        CREATE VIRTUAL TABLE foods_fts USING fts5(description, content='foods', content_rowid='id');
        INSERT INTO foods VALUES (1, 'Bananas, raw', 89, 1.1, 0.3, 22.8);
        INSERT INTO foods VALUES (2, 'Banana chips', 519, 2.3, 33.6, 58.4);
        INSERT INTO foods VALUES (3, 'Buckwheat groats, roasted, cooked', 92, 3.4, 0.6, 19.9);
        INSERT INTO foods_fts(foods_fts) VALUES ('rebuild');
    """)
    db.commit()
    return p


def test_lookup_order(st, usda):
    nu = Nutrition(st, usda, use_off=False)
    assert nu.lookup("банан", "banana raw").match == "USDA: Bananas, raw"
    # уточнение, которого нет в базе, отбрасывается
    assert nu.lookup("гречка", "buckwheat boiled").kcal == 92
    st.put_food("Банан", 95)
    assert nu.lookup("банан", "banana raw").match == "мой справочник"
    assert nu.lookup("чак-чак", "chak chak") is None
    assert nu.lookup("л-карнитин", "l carnitine syrup") is None  # не «Bananas» по обрывку «l»…
    assert nu.lookup("л-карнитин", "l carnitine syrup", guess=15).match == "≈ оценка модели"


def test_draft_totals(st, usda):
    d = Nutrition(st, usda, use_off=False).fill(Draft([Item("банан", "banana raw", 120), Item("чак-чак", "x", 50)]))
    assert round(d.total()) == 107
    assert [i.name for i in d.missing()] == ["чак-чак"]


def test_llm_json():
    d = to_draft('{"items":[{"name_ru":"Овсянка","name_en":"oats dry","grams":60},'
                 '{"name_ru":"","name_en":"x","grams":10}]}', "text")
    assert [(i.name, i.grams) for i in d.items] == [("овсянка", 60)]
    with pytest.raises(ValueError):
        to_draft('{"items":[]}', "text")
    d = to_draft('{"items":[{"name_ru":"сироп","name_en":"syrup","grams":25,"kcal_100g":30,"user_kcal_100g":14}]}', "text")
    assert d.items[0].per100.kcal == 14 and round(d.items[0].kcal, 1) == 3.5


# --- итоги и xlsx ---

def test_day_remaining_with_workout(st, tmp_path):
    day = date(2026, 10, 6)
    at = datetime(2026, 10, 6, 13, 0)
    st.add_entry("quick", "рацион", 1100, "manual", when=at)
    st.add_entry("workout", "бег", 200, "manual", when=at)
    st.add_entry("meal", "банан 120г", 107, "text", 1.3, 0.4, 27.4,
                 [{"name": "банан", "grams": 120, "kcal": 107}], when=at)
    assert report.goal(st, day) == 1350
    assert report.remaining(st, day) == 1350 - 1207 + 200
    assert "осталось 343" in report.format_day(st, day)

    st.set("goal_manual", 1000)
    assert report.left_line(st, day) == "перебор 7 ккал"

    st.add_weight(57.5, day)
    wb = load_workbook(report.build_xlsx(st, tmp_path / "x.xlsx"))
    assert wb.sheetnames == ["Журнал", "Дни", "Вес"]
    assert wb["Журнал"].max_row == 4
    assert wb["Журнал"]["E3"].value == -200


def test_week_low_average_note(st):
    for d in range(1, 8):
        st.add_entry("quick", "рацион", 1100, "manual", when=datetime(2026, 10, d, 12))
    text = report.format_week(st, date(2026, 10, 7))
    assert "в среднем 1 100" in text and "ниже 1200" in text


def test_dishes(st):
    st.put_dish("Цезарь  Жан-Жак", 450)
    assert st.get_dish("цезарь жан-жак")["kcal"] == 450
    assert st.get_dish("цезарь") is None
    st.put_dish("цезарь жан-жак", 520)
    assert [(r["name"], r["kcal"]) for r in st.dishes()] == [("цезарь жан-жак", 520)]
    assert st.delete_dish("ЦЕЗАРЬ жан-жак") and not st.dishes()


def test_day_starts_at_4am(tmp_path):
    path = str(tmp_path / "d.db")
    old = Storage(path, day_start_hour=0)  # как было раньше: день по календарю
    old.add_entry("quick", "лимонад", 100, "manual", when=datetime(2026, 10, 8, 0, 18))
    assert old.entries(date(2026, 10, 8))

    st = Storage(path, day_start_hour=4)  # старые записи пересчитываются при старте
    assert st.entries(date(2026, 10, 7))[0]["descr"] == "лимонад"
    st.add_entry("quick", "завтрак", 300, "manual", when=datetime(2026, 10, 8, 4, 0))
    assert st.day_totals(date(2026, 10, 8))["eaten"] == 300


def test_edit_delete_restore(st):
    eid = st.add_entry("meal", "овсянка 60г", 228, "text", 8, 4, 40,
                       [{"name": "овсянка", "grams": 60, "kcal": 228}])
    st.set_entry_kcal(eid, 114)  # полпорции — БЖУ тоже пополам
    e = st.get_entry(eid)
    assert (e["kcal"], e["protein"], e["fat"], e["carbs"]) == (114, 4, 2, 20)
    saved = st.delete_entry(eid)
    assert st.get_entry(eid) is None
    st.restore_entry(saved)
    assert st.get_entry(eid)["kcal"] == 114
    assert st.db.execute("SELECT COUNT(*) FROM entry_items WHERE entry_id = ?", (eid,)).fetchone()[0] == 1


def test_watch_day(st):
    day = date(2026, 10, 9)
    at = datetime(2026, 10, 9, 21, 30)
    st.add_entry("quick", "рацион", 1300, "manual", when=at)
    st.add_entry("workout", "прогулка", 150, "manual", when=at)
    assert st.day_totals(day)["burned"] == 150
    st.set_watch(412, when=at)          # прогулка внутри: засчитываем 412 − 200, а не 150 + 212
    assert st.day_totals(day)["burned"] == 212
    st.set_watch(300, when=at)          # новая цифра за день заменяет старую; ручные 150 больше
    t = st.day_totals(day)
    assert (t["burned"], t["eaten"]) == (150, 1300)
    assert len([e for e in st.entries(day) if e["kind"] == "watch"]) == 1
    assert "⌚ часы: 300 активных, сверх обычных 100" in report.format_day(st, day)


def test_water(st):
    day = date(2026, 10, 9)
    st.add_water(250, datetime(2026, 10, 9, 10))
    st.add_water(500, datetime(2026, 10, 9, 13))
    st.add_water(250, datetime(2026, 10, 10, 2))  # до 4 утра — ещё вчера
    assert st.water_on(day) == 1000
    assert st.remove_last_water(day) == 250 and st.water_on(day) == 750
    assert "💧 0,75 л" in report.format_day(st, day)


def test_undo(st):
    st.add_entry("quick", "a", 500, "manual")
    st.add_entry("meal", "b", 300, "text", items=[{"name": "b", "grams": 100, "kcal": 300}])
    assert st.delete_last()["descr"] == "b"
    assert st.db.execute("SELECT COUNT(*) FROM entry_items").fetchone()[0] == 0


def test_week_chart_png(st):
    st.add_entry("quick", "рацион", 1400, "manual", when=datetime(2026, 10, 8, 12))
    png = report.week_chart(st, date(2026, 10, 9))
    assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 10_000
