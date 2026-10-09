"""Тексты сводок и xlsx."""
import html
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import profile_calc
from nutrition import Draft
from storage import Storage

LOW_AVG_KCAL = 1200
WEEKDAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]


def goal(st: Storage, today: date | None = None) -> int:
    manual = st.get("goal_manual")
    if manual:
        return int(manual)
    today = today or date.today()
    return profile_calc.target(
        weight_kg=st.last_weight() or float(st.get("start_weight", 60)),
        height_cm=float(st.get("height_cm", 165)),
        age=today.year - int(st.get("birth_year", 1996)),
        sex=st.get("sex", "f"),
        activity=float(st.get("activity", 1.2)),
        deficit=float(st.get("deficit", 0.15)))


def remaining(st: Storage, day: date) -> int:
    t = st.day_totals(day)
    return round(goal(st, day) - t["eaten"] + t["burned"])


def esc(s: str) -> str:
    """Текст из баз и от пользователя -> безопасный для Telegram HTML."""
    return html.escape(html.unescape(str(s)), quote=False)


def n(x: float) -> str:
    return f"{round(x):,}".replace(",", " ")


def left_line(st: Storage, day: date) -> str:
    left = remaining(st, day)
    return f"осталось {n(left)} ккал" if left >= 0 else f"перебор {n(-left)} ккал"


def format_draft(d: Draft, st: Storage, day: date) -> str:
    lines = []
    for i in d.items:
        if i.per100:
            lines.append(f"• {esc(i.name)} {n(i.grams)} г · {n(i.kcal)} ккал  <i>({esc(i.per100.match)})</i>")
        else:
            lines.append(f"• {esc(i.name)} {n(i.grams)} г · ❓")
    total = d.total()
    pfc = f"бжу {n(d.total('protein'))}/{n(d.total('fat'))}/{n(d.total('carbs'))}"
    after = remaining(st, day) - total
    tail = f"останется {n(after)}" if after >= 0 else f"перебор {n(-after)}"
    return "\n".join(lines) + f"\n\n<b>итого {n(total)} ккал</b> · {pfc}\nпосле этого {tail}"


def format_day(st: Storage, day: date, title: str = "сегодня") -> str:
    t = st.day_totals(day)
    g = goal(st, day)
    rows = []
    for e in st.entries(day):
        if e["kind"] == "watch":
            extra = max(e["kcal"] - st.active_baseline, 0)
            rows.append(f"⌚ часы: {n(e['kcal'])} активных, сверх обычных {n(extra)}")
            continue
        sign = "−" if e["kind"] == "workout" else ""
        rows.append(f"{e['ts'][11:16]}  {sign}{n(e['kcal'])}  {esc(e['descr'])}")
    body = "\n".join(rows) or "пока пусто"
    return (f"<b>{title}</b>\n{body}\n\n"
            f"съедено {n(t['eaten'])} · тренировки +{n(t['burned'])} · норма {n(g)}\n"
            f"бжу {n(t['protein'])}/{n(t['fat'])}/{n(t['carbs'])}{water_line(st, day)}\n"
            f"<b>{left_line(st, day)}</b>")


def water_line(st: Storage, day: date) -> str:
    ml = st.water_on(day)
    return f" · 💧 {ml / 1000:.2f} л".replace(".", ",") if ml else ""


def format_week(st: Storage, end: date) -> str:
    start = end - timedelta(days=6)
    lines, eaten_days = [], []
    for k in range(7):
        d = start + timedelta(days=k)
        t = st.day_totals(d)
        if not t["n"]:
            lines.append(f"{WEEKDAYS[d.weekday()]} {d:%d.%m}  ·")
            continue
        left = remaining(st, d)
        mark = "✅" if left >= 0 else "🔺"
        burn = f" +{n(t['burned'])}" if t["burned"] else ""
        lines.append(f"{WEEKDAYS[d.weekday()]} {d:%d.%m}  {n(t['eaten'])}{burn}  {mark} {n(left)}")
        if t["eaten"]:
            eaten_days.append(t["eaten"])
    text = f"<b>неделя {start:%d.%m}-{end:%d.%m}</b>\n<pre>" + "\n".join(lines) + "</pre>"
    if eaten_days:
        avg = sum(eaten_days) / len(eaten_days)
        text += f"\nв среднем {n(avg)} ккал в день при норме {n(goal(st, end))}"
        if avg < LOW_AVG_KCAL:
            text += (f"\n\n💛 средняя ниже {LOW_AVG_KCAL}, это уже жёсткий дефицит. "
                     "если так неделями, стоит добавить немного еды")
    w = st.weights()
    if len(w) >= 2:
        text += f"\nвес: {w[0]['kg']} → {w[-1]['kg']} кг"
    return text


# цвета из эталонной палитры dataviz: одна серия — slot 1, текст — токены текста, не цвет серии
C_SURFACE, C_TEXT, C_MUTED, C_GRID, C_SERIES = "#fcfcfb", "#0b0b0b", "#52514e", "#e6e5e0", "#2a78d6"


def week_chart(st: Storage, end: date) -> bytes:
    """PNG: съедено по дням (столбцы) против «можно сегодня» = норма + активность (пунктир).
    Вес — отдельной панелью снизу, если записей хотя бы две (никаких двух шкал на одном графике)."""
    import io
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates
    import matplotlib.pyplot as plt

    days = [end - timedelta(days=6 - k) for k in range(7)]
    totals = [st.day_totals(d) for d in days]
    eaten = [t["eaten"] for t in totals]
    allowed = [goal(st, d) + t["burned"] for d, t in zip(days, totals)]
    weights = [w for w in st.weights() if days[0] - timedelta(days=21) <= date.fromisoformat(w["day"]) <= end]

    rows = 2 if len(weights) >= 2 else 1
    fig, axes = plt.subplots(rows, 1, figsize=(7, 4.2 if rows == 1 else 6.2), dpi=150,
                             gridspec_kw={"height_ratios": [3, 1.4][:rows]}, facecolor=C_SURFACE)
    axes = [axes] if rows == 1 else list(axes)
    ax = axes[0]
    xs = list(range(7))
    ax.bar(xs, eaten, width=0.56, color=C_SERIES, zorder=2)
    ax.step([-0.5, *xs, 6.5], [allowed[0], *allowed, allowed[-1]], where="mid",
            color=C_MUTED, linewidth=2, linestyle=(0, (4, 3)), zorder=3)
    for x, e, a in zip(xs, eaten, allowed):
        if e > a:
            ax.text(x, e + 25, f"+{n(e - a)}", ha="center", va="bottom", fontsize=9, color=C_TEXT)
    ax.text(6.55, allowed[-1], "можно\nс активностью", va="center", ha="left", fontsize=8, color=C_MUTED)
    ax.set_xticks(xs, [f"{WEEKDAYS[d.weekday()]}\n{d:%d.%m}" for d in days])
    days_with_food = sum(1 for e in eaten if e) or 1
    ax.set_title(f"съедено за неделю, ккал · в среднем {n(sum(eaten) / days_with_food)}",
                 loc="left", fontsize=11, color=C_TEXT)

    if rows == 2:
        wa = axes[1]
        wd = [date.fromisoformat(w["day"]) for w in weights]
        wa.plot(wd, [w["kg"] for w in weights], color=C_SERIES, linewidth=2, marker="o", markersize=5, zorder=2)
        wa.set_title("вес, кг", loc="left", fontsize=10, color=C_TEXT)
        wa.xaxis.set_major_formatter(matplotlib.dates.DateFormatter("%d.%m"))
        wa.text(wd[-1], weights[-1]["kg"], f"  {weights[-1]['kg']}", va="center", fontsize=9, color=C_TEXT)

    for a in axes:
        a.set_facecolor(C_SURFACE)
        a.grid(axis="y", color=C_GRID, linewidth=0.8, zorder=0)
        a.tick_params(colors=C_MUTED, labelsize=8, length=0)
        for side in ("top", "right", "left"):
            a.spines[side].set_visible(False)
        a.spines["bottom"].set_color(C_GRID)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=C_SURFACE)
    plt.close(fig)
    return buf.getvalue()


def build_xlsx(st: Storage, path: Path) -> Path:
    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Журнал"
    ws.append(["Дата", "Время", "Тип", "Описание", "Ккал", "Белки", "Жиры", "Углеводы", "Источник"])
    kinds = {"meal": "еда", "quick": "еда", "workout": "тренировка", "watch": "часы за день"}
    for e in st.all_entries():
        kcal = -e["kcal"] if e["kind"] in ("workout", "watch") else e["kcal"]
        ws.append([e["day"], e["ts"][11:16], kinds[e["kind"]], e["descr"], round(kcal),
                   *(round(e[k], 1) if e[k] is not None else None for k in ("protein", "fat", "carbs")),
                   e["source"]])

    days = sorted({e["day"] for e in st.all_entries()})
    ws2 = wb.create_sheet("Дни")
    ws2.append(["Дата", "Съедено", "Тренировки", "Норма", "Остаток", "Белки", "Жиры", "Углеводы"])
    for d in days:
        dd = date.fromisoformat(d)
        t = st.day_totals(dd)
        ws2.append([d, round(t["eaten"]), round(t["burned"]), goal(st, dd), remaining(st, dd),
                    round(t["protein"]), round(t["fat"]), round(t["carbs"])])

    ws3 = wb.create_sheet("Вес")
    ws3.append(["Дата", "Вес, кг"])
    for w in st.weights():
        ws3.append([w["day"], w["kg"]])

    for sheet in wb.worksheets:
        for c in sheet[1]:
            c.font = bold
        sheet.column_dimensions["A"].width = 12
    ws.column_dimensions["D"].width = 40
    wb.save(path)
    return path
